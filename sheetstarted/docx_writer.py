from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .organizer import WorksheetDraft


def _add_header_footer(document: Document, title: str) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = title
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer = section.footer.paragraphs[0]
    footer.text = "ชื่อ-สกุล: ____________________   คะแนน: ______"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_student_info_table(document: Document) -> None:
    table = document.add_table(rows=2, cols=3)
    headers = ["ชื่อ-สกุล", "ชั้น/ห้อง", "เลขที่"]
    for idx, value in enumerate(headers):
        table.cell(0, idx).text = value
        table.cell(1, idx).text = "________________________"


def write_docx(draft: WorksheetDraft, output_path: Path) -> Path:
    document = Document()
    document.add_heading("ใบงาน", level=0)
    document.add_paragraph(f"รายวิชา: {draft.subject}")
    document.add_paragraph(f"เรื่อง: {draft.lesson_title}")
    document.add_paragraph(f"ระดับชั้น: {draft.grade_level}")

    _add_header_footer(document, f"ใบงาน: {draft.lesson_title}")
    _add_student_info_table(document)

    for section in draft.sections:
        document.add_heading(section.title, level=1)
        for item in section.content:
            document.add_paragraph(item)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
